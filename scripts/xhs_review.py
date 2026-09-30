"""Offline Xiaohongshu evidence scoring and versioned workbook export. No network/send."""
import argparse, copy, json, math, re
from pathlib import Path
from statistics import median
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import ColorScaleRule

WEIGHTS={'score_fit':30,'score_engagement':25,'score_audience':20,'score_content':15,'score_value':10}
HUMAN={'status','approval','first_contact','last_contact','private_image','negotiated','planned_budget','communication','outreach_text'}
HEADERS=['编号','平台','博主昵称','垂类','粉丝数(万)','近30天更新条数','中位赞藏/播放（注明口径）','匹配打分(0-100)','量级','联系渠道','触达状态','首次触达日期','最近跟进日期','私聊报价-图文(元)','蒲公英刊例-图文(元)','刊例报价-视频(元)','同格式合作阅读参考（非承诺）','建议合作形式','拟投预算(用户确认)','千次阅读成本（预算口径）','意向度','老板审批','沟通记录/备注','垂类匹配度','互动质量','粉丝画像','内容质量','商单性价比','主页链接','数据采集备注','数据可信度','梯队','梯队内排名','评分状态','评分覆盖率','匹配分依据','合作样本数','合作格式/窗口/流量','最新谈价(元)','谈价格式/含费','建议询价净价(非报价)','建议目标/动作','平台预估曝光CPM','来源种子/发现路径','小红书号','后续需核实']

def num(v):
    if isinstance(v,bool):return None
    try:
        f=float(v)
        return f if math.isfinite(f) else None
    except (ValueError,TypeError):return None

def score_creator(c):
    if c.get('stage') not in ('截图初评','内容核实通过') or c.get('excluded'):return {'total':None,'coverage':0,'values':{},'label':'待核实' if not c.get('excluded') else '已排除'}
    values={}
    for key,w in WEIGHTS.items():
        e=c.get('scores',{}).get(key,{})
        v=num(e.get('value'))
        if v is None:continue
        if not 0<=v<=100:raise ValueError('score out of bounds')
        if not e.get('reason') or not e.get('source'):raise ValueError('score requires evidence')
        if key!='score_fit' and e.get('level')!='verified':continue
        if key=='score_audience' and not e.get('official_audience'):continue
        if key=='score_fit' and e.get('parts'):
            parts=e['parts'];caps=[40,25,20,15]
            if len(parts)!=4 or any(num(p) is None or p<0 or p>cap for p,cap in zip(parts,caps)) or sum(parts)!=v:raise ValueError('fit parts mismatch')
        values[key]=v
    den=sum(WEIGHTS[k] for k in values)
    # A rich commercial dataset alone cannot create a content-match score.
    if 'score_fit' not in values:return {'total':None,'coverage':den/100,'values':values,'label':'缺匹配证据'}
    return {'total':round(sum(values[k]*WEIGHTS[k] for k in values)/den,2),'coverage':den/100,'values':values,'label':('截图初评·非完整排名' if c['stage']=='截图初评' else '核验分')+('·部分维度' if den<100 else '')}

def cost_per_thousand(cost,denominator):
    c,d=num(cost),num(denominator)
    return round(c/d*1000,2) if c is not None and c>=0 and d is not None and d>0 else None

def comparable_reference(c):
    m=c.get('commercial',{})
    return (m.get('read_median') if m.get('format_confirmed') and m.get('format')==c.get('format') and m.get('sample_n') and m.get('kind')=='合作' else None)

def mature_median(notes,fmt):
    seen=set();values=[]
    for n in notes:
        key=n.get('note_id') or (n.get('date'),n.get('title'))
        if key in seen:continue
        seen.add(key)
        if n.get('kind')=='合作' and n.get('format')==fmt and n.get('age_days',-1)>=7 and num(n.get('reads')) is not None:values.append(n['reads'])
    return median(values) if values else None,len(values)

def merge_creator(old,new):
    if old['id']!=new['id']:raise ValueError('identity mismatch')
    merged=copy.deepcopy(old);conflicts=[]
    for k,v in new.items():
        if k in HUMAN and old.get(k) not in (None,'') and v!=old[k]:conflicts.append(k);continue
        merged[k]=v
    if old.get('excluded'):
        merged['excluded']=True;merged['status']='已否决'
    return merged,conflicts

def export_outreach(data,path):
    path=Path(path)
    if path.exists():raise FileExistsError('Use a new output version')
    parts=['# 第二轮询价草稿\n\n未发送。历史联系者发送前核对最新聊天；建议价不是成交价。主页未核实的条目不可作为发送定位依据。']
    for c in data['creators']:
        if not c.get('outreach_text') or c.get('excluded'):continue
        url=c.get('profile_url')
        link=f'[打开主页]({url})' if url and url.startswith('https://www.xiaohongshu.com/user/profile/') else '待核实'+('（小红书号：'+str(c['account'])+'）' if c.get('account') else '')
        parts.append(f"## {c['name']}\n\n主页链接：{link}\n\n状态：草稿，未发送\n\n私信文案：\n\n{c['outreach_text']}")
    path.write_text('\n\n'.join(parts),encoding='utf-8')

def export(data,path,outreach=None):
    path=Path(path)
    if path.exists():raise FileExistsError('Use a new output version; never overwrite human workbook')
    ids=[c['id'] for c in data['creators']]
    if len(set(ids))!=len(ids):raise ValueError('duplicate identity')
    w=Workbook();info=w.active;info.title='1-使用说明'
    info.append(['项目','说明'])
    for k,v in data.get('notes',{}).items():info.append([k,str(v)])
    standards=w.create_sheet('2-筛选标准与打法');standards.append(['环节','规则'])
    for row in [
        ('统一候选池','所有类型在3-博主库；类型只是标签，非分池。列表发现≠内容已核实。'),
        ('匹配评分','五维权重30/25/20/15/10，按有效权重归一化；无画像不填默认分；评分覆盖率必须同行查看。'),
        ('截图初评','本轮已评分者只填可观察的匹配分；综合分与匹配分相等因仅30%权重可用，不代表其余维度合格。'),
        ('匹配分四项','任务相关40＋重复性25＋具体信息20＋可演示性15；每项理由见证据页；截图不能确认正文/亲历。'),
        ('完整核验','90天最近10篇非置顶，打开至少3篇相关正文；记录真实任务和约束；不足不凑数。'),
        ('相似迭代','核实种子→相似列表→10人复核→合格新种子→最多2层；记录来源，弱分支暂停。'),
        ('合作口径','同格式、同窗口、已成熟合作作参考；混合/样本未知不自动填Q。推广全流量并非自然量。'),
        ('费用','正式预算仅保留用户原值；内部情景见5-报价与成本，暂定10%必须另标；不是平台实际费率。'),
        ('人工字段','未来合并保护报价、谈价、预算、审批、跟进；冲突提示，禁止静默覆盖；输出新版本。'),
        ('询价','不设询价sheet；可选独立MD包含名字、主页链接、私信文案，不自动发送。'),
        ('范围','仅小红书；抖音、海外规则与脚本未修改。')]:standards.append(row)
    main=w.create_sheet('3-博主库');main.append(HEADERS)
    evidence=w.create_sheet('4-内容与商单证据');evidence.append(['博主','证据类型','指标/标题/维度','值/内容','日期/窗口','格式/口径','来源','限制/判断'])
    costs=w.create_sheet('5-报价与成本');costs.append(['博主','拟购格式','刊例图文','刊例视频','私聊历史','最新谈价','报价/费用说明','建议询价(非报价)','目标','观察合作阅读','样本数','格式/窗口/流量','平台预估曝光CPM','假设费用率','刊例净价情景','情景总成本','千次阅读情景成本','是否可比','限制'])
    history=w.create_sheet('6-排除与历史');history.append(['类型','原编号','博主','原状态','原审批','原因/历史数据'])
    for r in data.get('history',[]):history.append(r)
    for i,c in enumerate(data['creators'],2):
        if c.get('excluded'):raise ValueError('Excluded creators must be in history, not active pool')
        s=score_creator(c);v=s['values'];m=c.get('commercial',{});ref=comparable_reference(c)
        f=c.get('followers');tier=('KOC' if f<10000 else '初级' if f<50000 else '腰部及以上') if isinstance(f,(int,float)) else ''
        sc=c.get('scores',{}).get('score_fit',{})
        row=[c['id'],'小红书',c['name'],c.get('type','待核实'),f/10000 if f else None,c.get('posts_30d'),None,s['total'],tier,'小红书/蒲公英',c.get('status','待核实'),c.get('first_contact'),c.get('last_contact'),c.get('private_image'),c.get('quote_image'),c.get('quote_video'),ref,c.get('format'),c.get('planned_budget'),None,c.get('priority'),c.get('approval'),c.get('communication'),*[v.get(k) for k in WEIGHTS],c.get('profile_url'),c.get('source'),c.get('stage'),None,None,s['label'],s['coverage'],sc.get('reason'),m.get('sample_n'),m.get('scope'),c.get('negotiated'),c.get('negotiated_scope'),c.get('ask'),c.get('target'),c.get('platform_cpm'),c.get('discovery'),c.get('account'),c.get('gaps')]
        main.append(row)
        # Formula weighted terms only refer to scores that passed evidence validation.
        if s['total'] is not None:
            terms=[];dens=[]
            for col,key in zip(['X','Y','Z','AA','AB'],WEIGHTS):
                terms.append(f'N({col}{i})*{WEIGHTS[key]}');dens.append(f'IF(ISNUMBER({col}{i}),{WEIGHTS[key]},0)')
            main.cell(i,8,'=IFERROR(ROUND(('+ '+'.join(terms)+')/('+ '+'.join(dens)+'),2),"")')
        if ref and c.get('budget_scope_verified'):
            main.cell(i,20,f'=IF(AND(ISNUMBER(S{i}),S{i}>0,ISNUMBER(Q{i}),Q{i}>0),ROUND(S{i}/Q{i}*1000,2),"")')
        if c.get('profile_url'):main.cell(i,29).hyperlink=c['profile_url'];main.cell(i,29).style='Hyperlink'
        main.cell(i,35).number_format='0%'
        for key,se in c.get('scores',{}).items():evidence.append([c['name'],'评分证据',key,str(se.get('parts') or se.get('value')),data.get('as_of'),'截图初评' if se.get('level')!='verified' else '已核实',se.get('source'),se.get('reason')])
        if m:evidence.append([c['name'],'合作统计','阅读中位数',m.get('read_median'),data.get('as_of'),m.get('scope'),c.get('source'),m.get('limitation')])
        price=c.get('quote_image') if c.get('format')=='图文' else c.get('quote_video') if c.get('format')=='视频' else None
        total=round(price*1.1,2) if price is not None else None
        costs.append([c['name'],c.get('format'),c.get('quote_image'),c.get('quote_video'),c.get('historical_quotes'),c.get('negotiated'),c.get('negotiated_scope'),c.get('ask'),c.get('target'),m.get('read_median'),m.get('sample_n'),m.get('scope'),c.get('platform_cpm'),.1 if price else None,price,total,cost_per_thousand(total,ref),'同格式观察参考' if ref else '不可作为拟购格式成本',m.get('limitation','')+'；假设另加10%，不含未知推广费用；非市场报价或实际效果'])
    for r in data.get('evidence',[]):evidence.append(r)
    for sh in w:
        sh.freeze_panes='D2' if sh==main else 'B2';sh.auto_filter.ref=sh.dimensions
        sh.sheet_view.zoomScale=80
        for cell in sh[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='193B50');cell.alignment=Alignment(wrap_text=True,vertical='center')
        sh.row_dimensions[1].height=48
        for col in sh.columns:sh.column_dimensions[col[0].column_letter].width=19
        for row in sh.iter_rows(min_row=2):
            sh.row_dimensions[row[0].row].height=64
            for cell in row:
                cell.alignment=Alignment(wrap_text=True,vertical='top')
                if cell.row%2==0:cell.fill=PatternFill('solid',fgColor='F0F5F8')
    for sh in [info,standards]:sh.column_dimensions['B'].width=115
    for col in ['W','AD','AJ','AL','AP','AR','AT']:main.column_dimensions[col].width=45
    main.column_dimensions['C'].width=25;main.column_dimensions['AC'].width=45
    for col in ['D','F','G','H']:evidence.column_dimensions[col].width=50
    costs.column_dimensions['L'].width=55;costs.column_dimensions['S'].width=65
    history.column_dimensions['F'].width=100
    if main.max_row>1:
        dv=DataValidation(type='list',formula1='"待核实,待触达,已私信,已回复,已报价,已合作,已否决,暂停"');main.add_data_validation(dv);dv.add(f'K2:K{main.max_row}')
        main.conditional_formatting.add(f'X2:X{main.max_row}',ColorScaleRule(start_type='num',start_value=0,start_color='FBE5D6',end_type='num',end_value=100,end_color='C6E0B4'))
    w.save(path)
    # Persist formula caches for viewers that do not calculate Excel formulas.
    import zipfile,xml.etree.ElementTree as ET,os
    ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    with zipfile.ZipFile(path) as z:files={n:z.read(n) for n in z.namelist()}
    xml=ET.fromstring(files['xl/worksheets/sheet3.xml'])
    cache={}
    for i,c in enumerate(data['creators'],2):
        cache[f'H{i}']=score_creator(c)['total'];cache[f'T{i}']=cost_per_thousand(c.get('planned_budget'),comparable_reference(c)) if c.get('budget_scope_verified') else None
    for cell in xml.findall('.//m:c',ns):
        if cell.find('m:f',ns) is not None:
            val=cell.find('m:v',ns)
            if val is None:val=ET.SubElement(cell,'{'+ns['m']+'}v')
            n=cache.get(cell.attrib['r']);val.text=str(n) if n is not None else None
    files['xl/worksheets/sheet3.xml']=ET.tostring(xml,encoding='utf-8',xml_declaration=True)
    temp=path.with_suffix('.tmp.xlsx')
    with zipfile.ZipFile(temp,'w',zipfile.ZIP_DEFLATED) as z:
        for n,b in files.items():z.writestr(n,b)
    os.replace(temp,path)
    if outreach:export_outreach(data,outreach)
    return {'creators':len(data['creators']),'scored':sum(score_creator(c)['total'] is not None for c in data['creators']),'sheets':w.sheetnames}

if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('command',choices=['export']);p.add_argument('data');p.add_argument('output');p.add_argument('--outreach');a=p.parse_args()
    print(json.dumps(export(json.loads(Path(a.data).read_text(encoding='utf-8')),a.output,a.outreach),ensure_ascii=False))
